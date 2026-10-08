import pytest
import openpyxl
from io import BytesIO
from app.services.doctor_service import DoctorService


@pytest.mark.asyncio
async def test_doctor_bulk_template_generation_and_styling():
    service = DoctorService(None)
    stream = await service.generate_doctor_bulk_template()
    assert stream is not None

    wb = openpyxl.load_workbook(stream)
    ws = wb.active
    assert ws.title == "Doctor Bulk Template"

    header_row = [cell.value for cell in ws[1]]
    expected_mandatory_headers = [
        "First Name *", "Last Name *", "Specialization *", "Experience *",
        "Phone *", "Email *", "Password *", "Department Name *",
        "Consultation Fee *", "License Number *", "Gender *"
    ]
    expected_optional_headers = [
        "Qualification", "Availability Status", "Bio", "Date of Birth"
    ]

    for header in expected_mandatory_headers:
        assert header in header_row, f"Mandatory header {header} missing from template"

    for header in expected_optional_headers:
        assert header in header_row, f"Optional header {header} missing from template"

    # Verify cell styles
    for cell in ws[1]:
        if "*" in str(cell.value):
            assert cell.fill.start_color.rgb in ("00FFC7CE", "FFC7CE")
            assert cell.font.bold is True
            assert cell.font.color.rgb in ("009C0006", "9C0006")
        else:
            assert cell.fill.start_color.rgb in ("00DCE6F1", "DCE6F1")
            assert cell.font.bold is True
            assert cell.font.color.rgb in ("001F497D", "1F497D")


def test_doctor_bulk_header_normalization_with_asterisks():
    raw_headers = [
        "First Name *", "Last Name *", "Specialization *", "Qualification", "Experience *",
        "Phone *", "Email *", "Password *", "Department Name *", "Consultation Fee *",
        "License Number *", "Availability Status", "Bio", "Gender *", "Date of Birth"
    ]
    normalized_headers = [str(h).replace("*", "").strip().lower() for h in raw_headers if h is not None]

    required_headers = {
        "first name", "last name", "specialization", "experience",
        "phone", "email", "password", "license number",
        "department name", "consultation fee", "gender"
    }

    assert required_headers.issubset(set(normalized_headers))
    assert "first name" in normalized_headers
    assert "department name" in normalized_headers
    assert "consultation fee" in normalized_headers
    assert "gender" in normalized_headers
