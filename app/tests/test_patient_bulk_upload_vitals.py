import pytest
import openpyxl
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock
from datetime import date
from pydantic import ValidationError

from app.schemas.patient_schema import PatientCreate, PatientUpdate, PatientResponse
from app.services.patient_service import PatientService
from app.models.patient_model import Patient


def test_patient_create_schema_valid_weight_and_bp():
    data = {
        "first_name": "Rohan",
        "last_name": "Sharma",
        "gender": "Male",
        "dob": date(1990, 1, 1),
        "blood_group": "B+",
        "diagnosis": "Healthy checkup",
        "weight": 72.5,
        "blood_pressure": "120/80",
    }
    patient_create = PatientCreate(**data)
    assert patient_create.weight == 72.5
    assert patient_create.blood_pressure == "120/80"


def test_patient_create_schema_optional_weight_and_bp():
    data = {
        "first_name": "Priya",
        "last_name": "Patil",
        "gender": "Female",
        "diagnosis": "Consultation",
    }
    patient_create = PatientCreate(**data)
    assert patient_create.weight is None
    assert patient_create.blood_pressure is None


def test_patient_create_schema_invalid_weight():
    with pytest.raises(ValidationError) as exc:
        PatientCreate(
            first_name="Rohan",
            last_name="Sharma",
            diagnosis="Routine",
            weight=-5.0,
        )
    assert "Weight must be greater than 0" in str(exc.value)

    with pytest.raises(ValidationError) as exc2:
        PatientCreate(
            first_name="Rohan",
            last_name="Sharma",
            diagnosis="Routine",
            weight=600.0,
        )
    assert "Weight cannot exceed 500 kg" in str(exc2.value)


def test_patient_create_schema_invalid_bp():
    with pytest.raises(ValidationError) as exc:
        PatientCreate(
            first_name="Rohan",
            last_name="Sharma",
            diagnosis="Routine",
            blood_pressure="A" * 25,
        )
    assert "Blood pressure cannot exceed 20 characters" in str(exc.value)


@pytest.mark.asyncio
async def test_generate_patient_bulk_template_contains_weight_and_bp():
    mock_db = AsyncMock()
    service = PatientService(mock_db)

    stream = await service.generate_patient_bulk_template()
    wb = openpyxl.load_workbook(stream)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    assert "Weight (kg)" in headers
    assert "Blood Pressure" in headers

    # Verify sample data row
    sample_row = [cell.value for cell in ws[2]]
    weight_idx = headers.index("Weight (kg)")
    bp_idx = headers.index("Blood Pressure")

    assert sample_row[weight_idx] == 70.5
    assert sample_row[bp_idx] == "120/80"


@pytest.mark.asyncio
async def test_import_patients_from_excel_with_vitals():
    mock_db = AsyncMock()
    service = PatientService(mock_db)
    service.create = AsyncMock()

    # Create in-memory workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "First Name *", "Last Name *", "Gender", "Diagnosis *",
        "Weight (kg)", "Blood Pressure", "Phone", "Email"
    ])
    ws.append([
        "Amit", "Kumar", "Male", "Checkup",
        75.5, "130/85", "9876543210", "amit@example.com"
    ])
    ws.append([
        "Sunita", "Verma", "Female", "General",
        None, None, "9876543211", "sunita@example.com"
    ])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    mock_file = AsyncMock()
    mock_file.read.return_value = stream.getvalue()

    result = await service.import_patients_from_excel(
        file=mock_file,
        user_id=1,
    )

    assert result["total_rows"] == 2
    assert result["created"] == 2
    assert result["failed"] == 0
    assert result["errors"] == []

    assert service.create.call_count == 2
    # Verify first patient created with weight and blood_pressure
    first_call_arg = service.create.call_args_list[0][0][0]
    assert isinstance(first_call_arg, PatientCreate)
    assert first_call_arg.first_name == "Amit"
    assert first_call_arg.weight == 75.5
    assert first_call_arg.blood_pressure == "130/85"

    # Verify second patient created with None for weight and blood_pressure
    second_call_arg = service.create.call_args_list[1][0][0]
    assert isinstance(second_call_arg, PatientCreate)
    assert second_call_arg.first_name == "Sunita"
    assert second_call_arg.weight is None
    assert second_call_arg.blood_pressure is None


@pytest.mark.asyncio
async def test_import_patients_from_excel_invalid_weight():
    mock_db = AsyncMock()
    service = PatientService(mock_db)
    service.create = AsyncMock()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "First Name *", "Last Name *", "Gender", "Diagnosis *",
        "Weight (kg)", "Blood Pressure"
    ])
    ws.append([
        "Invalid", "Patient", "Male", "Checkup",
        "not_a_number", "120/80"
    ])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    mock_file = AsyncMock()
    mock_file.read.return_value = stream.getvalue()

    result = await service.import_patients_from_excel(
        file=mock_file,
        user_id=1,
    )

    assert result["total_rows"] == 1
    assert result["created"] == 0
    assert result["failed"] == 1
    assert len(result["errors"]) == 1
    assert "weight" in result["errors"][0]["error"].lower()

