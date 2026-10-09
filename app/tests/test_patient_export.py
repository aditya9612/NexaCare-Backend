import sys
import os
import asyncio
import uuid
import pytest
from httpx import AsyncClient, ASGITransport
from openpyxl import load_workbook
from io import BytesIO
from PIL import Image as PILImage
from sqlalchemy import select

if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

from app.main import app
from app.models.patient_model import Patient
from app.models.hospital_model import Hospital
from app.models.user_model import User
from app.core.database import AsyncSessionLocal, engine
from app.core.dependencies import get_current_active_user, get_current_user
from app.utils.helpers import utc_now


@pytest.fixture(autouse=True)
async def cleanup_engine():
    if sys.platform == "win32":
        asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())
    await engine.dispose()
    yield
    await engine.dispose()


def create_dummy_image_bytes():
    buf = BytesIO()
    img = PILImage.new("RGB", (100, 100), color="blue")
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_patients_export_pdf_and_excel_null_logo():
    if sys.platform == "win32":
        asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

    uid = uuid.uuid4().hex[:8]

    async with AsyncSessionLocal() as db:
        real_user = (await db.execute(select(User).limit(1))).scalar_one_or_none()
        real_user_id = real_user.id if real_user else 1

        hosp = Hospital(
            name=f"Test Hospital NullLogo {uid}",
            email=f"hospital_nulllogo_{uid}@test.com",
            phone="+919988776655",
            address="123 Care Street, Health City",
            logo_path=None,
        )
        db.add(hosp)
        await db.flush()
        
        p = Patient(
            first_name=f"Export_{uid}",
            last_name="Patient",
            patient_code=f"PAT_EXP_{uid}",
            gender="Female",
            status="active",
            hospital_id=hosp.id,
            phone="+919876543210",
            email=f"patient_{uid}@test.com",
            created_at=utc_now(),
        )
        db.add(p)
        await db.commit()
        hosp_id = hosp.id

    class FakeUser:
        id = real_user_id
        role_id = 1
        hospital_id = hosp_id
        is_active = True
        is_verified = True
        email = f"admin_{uid}@test.com"

        class role:
            name = "hospital_admin"

    fake_user = FakeUser()
    app.dependency_overrides[get_current_active_user] = lambda: fake_user
    app.dependency_overrides[get_current_user] = lambda: fake_user

    from unittest.mock import patch, AsyncMock
    auth_headers = {"Authorization": "Bearer fake-token"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test", headers=auth_headers) as ac:
        with patch("app.core.dependencies.require_permission") as mock_perm:
            mock_perm.return_value = lambda: fake_user
            with patch(
                "app.repositories.rbac_repository.RBACRepository.get_user_permissions",
                new=AsyncMock(return_value=["patients:read"]),
            ):
                # 1. Test Excel Export (Logo Null)
                resp_excel = await ac.get("/api/v1/patients/export?format=excel")
                assert resp_excel.status_code == 200
                assert resp_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                
                excel_wb = load_workbook(BytesIO(resp_excel.content))
                ws = excel_wb.active
                assert ws.cell(row=1, column=1).value == f"Test Hospital NullLogo {uid}"
                assert ws.freeze_panes == "A6"
                assert ws.cell(row=5, column=1).value == "Sr. No."

                # 2. Test PDF Export (Logo Null)
                resp_pdf = await ac.get("/api/v1/patients/export?format=pdf")
                assert resp_pdf.status_code == 200
                assert resp_pdf.headers["content-type"] == "application/pdf"
                assert len(resp_pdf.content) > 1000

    app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_hospital_logo_upload_and_export_with_logo():
    if sys.platform == "win32":
        asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

    uid = uuid.uuid4().hex[:8]

    async with AsyncSessionLocal() as db:
        real_user = (await db.execute(select(User).limit(1))).scalar_one_or_none()
        real_user_id = real_user.id if real_user else 1

        hosp = Hospital(
            name=f"Test Hospital WithLogo {uid}",
            email=f"hospital_withlogo_{uid}@test.com",
            phone="+919988776611",
            address="456 Health Ave, Wellness City",
            logo_path=None,
        )
        db.add(hosp)
        await db.flush()

        p = Patient(
            first_name=f"WithLogo_{uid}",
            last_name="Patient",
            patient_code=f"PAT_LOGO_{uid}",
            gender="Male",
            status="active",
            hospital_id=hosp.id,
            created_at=utc_now(),
        )
        db.add(p)
        await db.commit()
        hosp_id = hosp.id

    class FakeSuperAdmin:
        id = real_user_id
        role_id = 1
        hospital_id = hosp_id
        is_active = True
        is_verified = True
        email = f"superadmin_{uid}@test.com"

        class role:
            name = "super_admin"

    fake_admin = FakeSuperAdmin()
    app.dependency_overrides[get_current_active_user] = lambda: fake_admin
    app.dependency_overrides[get_current_user] = lambda: fake_admin

    from app.api.v1.routes.super_admin_routes import require_super_admin
    app.dependency_overrides[require_super_admin] = lambda: fake_admin

    from unittest.mock import patch, AsyncMock
    auth_headers = {"Authorization": "Bearer fake-token"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test", headers=auth_headers) as ac:
        with patch("app.core.dependencies.require_permission") as mock_perm:
            mock_perm.return_value = lambda: fake_admin
            with patch(
                "app.repositories.rbac_repository.RBACRepository.get_user_permissions",
                new=AsyncMock(return_value=["patients:read"]),
            ):
                # 1. Test Hospital Logo Upload API
                img_bytes = create_dummy_image_bytes()
                files = {"file": ("hospital_logo.png", img_bytes, "image/png")}
                resp_upload = await ac.post(f"/api/v1/super-admin/hospitals/{hosp_id}/logo", files=files)
                assert resp_upload.status_code == 200, f"Upload failed: {resp_upload.text}"
                data = resp_upload.json()["data"]
                saved_logo_path = data["logo_path"]
                assert saved_logo_path is not None
                assert os.path.exists(saved_logo_path)

                # 2. Test Excel Export with embedded logo
                resp_excel = await ac.get("/api/v1/patients/export?format=excel")
                assert resp_excel.status_code == 200
                excel_wb = load_workbook(BytesIO(resp_excel.content))
                ws = excel_wb.active
                assert ws.cell(row=1, column=1).value == f"Test Hospital WithLogo {uid}"
                assert len(ws._images) == 1, "Excel export must contain exactly 1 embedded logo image"

                # 3. Test PDF Export with logo
                resp_pdf = await ac.get("/api/v1/patients/export?format=pdf")
                assert resp_pdf.status_code == 200
                assert len(resp_pdf.content) > 1000

                # Clean up saved logo file
                if saved_logo_path and os.path.exists(saved_logo_path):
                    try:
                        os.remove(saved_logo_path)
                    except Exception:
                        pass

    app.dependency_overrides.clear()
